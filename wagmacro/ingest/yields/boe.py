"""GBP — Bank of England government liability curve, 2-year nominal spot.

The BoE publishes its curves as a zip of spreadsheets rather than a feed.
Verified on 2026-10-15, inside `GLC Nominal daily data current month.xlsx`,
sheet `4. spot curve`:

    years:      0.5      1       1.5     2       2.5     ...  5
    2026-10-01  4.1213  4.3500  4.5001  4.6008  4.6731  ... 4.9224

so the UK nominal 2-year spot was 4.6008 that day.

WHY NOT THE IADB. The Interactive Database is a far easier CSV endpoint
and it works — but its IUDSNPY returned 4.9288 on the same day, which
lines up with the FIVE-year column above (4.9224), not the two. Using it
would have put a number 33bp wrong into the most heavily weighted
component, every day, while looking entirely plausible.

ANCHORING. The maturity column is found by reading the row that starts
with "years:" and locating the entry equal to 2. It is deliberately NOT
found by scanning the first few rows for the value 2.0 — the sheet above
it has a "months:" row where 2.0 means two MONTHS. Anchoring on the
wrong row is how a parser quietly reports a 2-month rate as a 2-year one.

HISTORY. The published file covers the current month only, so early in a
month there are too few rows for the 20-observation lookback and the
scorer correctly degrades to "level only". A longer archive is attempted
first; its URL is unverified, and the diagnostic reports which file the
series actually came from.
"""

from __future__ import annotations

import io
import zipfile
from datetime import date, datetime

from ...core.types import Currency
from .base import ProviderResult, Series

#: Tried in order. The second is the verified one.
URLS = (
    "https://www.bankofengland.co.uk/-/media/boe/files/statistics/"
    "yield-curves/glcnominalddata.zip",
    "https://www.bankofengland.co.uk/-/media/boe/files/statistics/"
    "yield-curves/latest-yield-curve-data.zip",
)

MEMBER_MATCH = "glc nominal"
SHEET_MATCH = "spot curve"
YEARS_ROW_PREFIX = "years"
TENOR_YEARS = 2.0


def parse(blob: bytes) -> Series:
    """Pull the 2-year nominal spot series out of the archive."""
    try:
        archive = zipfile.ZipFile(io.BytesIO(blob))
    except zipfile.BadZipFile:
        return ()

    members = [n for n in archive.namelist()
               if MEMBER_MATCH in n.lower() and n.lower().endswith(".xlsx")]
    if not members:
        return ()

    try:
        import openpyxl
    except ImportError:                      # pragma: no cover - env issue
        return ()

    collected: dict[date, float] = {}
    for member in members:
        try:
            book = openpyxl.load_workbook(io.BytesIO(archive.read(member)),
                                          read_only=True, data_only=True)
        except Exception:                    # noqa: BLE001 - one bad file only
            continue
        for sheet_name in book.sheetnames:
            if SHEET_MATCH not in sheet_name.lower():
                continue
            collected.update(_read_sheet(book[sheet_name]))
        book.close()

    return tuple(sorted(collected.items()))


def _read_sheet(sheet) -> dict[date, float]:
    column: int | None = None
    out: dict[date, float] = {}

    for row in sheet.iter_rows(values_only=True):
        if not row:
            continue
        first = row[0]

        if column is None:
            # Only the "years:" row may define the column.
            if isinstance(first, str) and \
                    first.strip().lower().startswith(YEARS_ROW_PREFIX):
                for i, cell in enumerate(row):
                    if isinstance(cell, (int, float)) and \
                            abs(float(cell) - TENOR_YEARS) < 1e-6:
                        column = i
                        break
            continue

        if len(row) <= column:
            continue
        when = _as_date(first)
        value = row[column]
        if when is None or not isinstance(value, (int, float)):
            continue
        out[when] = float(value)

    return out


def _as_date(cell) -> date | None:
    if isinstance(cell, datetime):
        return cell.date()
    if isinstance(cell, date):
        return cell
    if isinstance(cell, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%b-%Y"):
            try:
                return datetime.strptime(cell.strip()[:10], fmt).date()
            except ValueError:
                continue
    return None


def fetch(get) -> ProviderResult:
    attempts: list[str] = []
    for target in URLS:
        status, body = get(target)
        name = target.rsplit("/", 1)[-1]
        if status != 200:
            attempts.append(f"{name}: HTTP {status}")
            continue
        series = parse(body)
        if series:
            return ProviderResult(
                Currency.GBP, target, series=series, ok=True,
                detail=f"{len(series)} obs from {name}, "
                       f"latest {series[-1][0]}")
        attempts.append(f"{name}: 200 but no 2y spot series found")
    return ProviderResult(Currency.GBP, URLS[-1],
                          detail="; ".join(attempts) or "no attempt made")
